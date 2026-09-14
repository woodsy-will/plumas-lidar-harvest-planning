"""06_review_sheets.py - simulated cruise, QA checks and one review sheet per unit.

A BAF 20 variable-radius cruise on a 300 ft grid is simulated inside each unit from the canopy products
(dominant height and cover drive the tree list). Six recording errors are planted on purpose and listed in
data/work/qc_planted.json. The QA pass flags plots that fail checks; the review sheet for each unit carries
stand metrics, the flagged plots and a plot map, in the form a crew lead hands back to the field.

Outputs data/work/cruise_plots.gpkg, output/review/Unit_<id>_Review.pdf, output/review/Unit_Reviews.pdf,
        output/review/qa_summary.csv and output/review/cruise_data.xlsx (sheets READ ME, Plots, Trees, Unit summary,
        Stand tables, Stock tables, Standards and assumptions). Every plot and tree is simulated: the review-sheet pages carry
        a SIMULATED DATA watermark, the cruiser code is SIM, the date field reads "simulated" and qa_summary.csv has a data column.
Run: python-qgis-ltr.bat scripts/06_review_sheets.py
"""
import json
import math
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from osgeo import gdal, ogr, osr
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
from scipy.ndimage import map_coordinates

gdal.UseExceptions(); ogr.UseExceptions(); osr.UseExceptions()
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW, WORK, OUT = (os.path.join(ROOT, p) for p in (os.path.join("data", "raw"), os.path.join("data", "work"), os.path.join("output", "review")))
os.makedirs(OUT, exist_ok=True)
BAF = 20.0; GRID = 300.0
MIN_PLOTS = 20         # R5 Supplement 2409.12-2021-4 sec. 41.3: "For area based (plot) cruises, there shall be a minimum of 20 plots"
TARGET_SDI_PCT = 35    # leave target at the lower limit of full site occupancy, 35 % of maximum SDI (Long 1985)
SPECIES = {"PP": ("Ponderosa pine", 0.30), "WF": ("White fir", 0.30), "DF": ("Douglas-fir", 0.15), "SP": ("Sugar pine", 0.10), "IC": ("Incense-cedar", 0.15)}
rng = np.random.default_rng(20260910)


def raster(name):
    d = gdal.Open(os.path.join(WORK, name)); b = d.GetRasterBand(1); a = b.ReadAsArray().astype("float32"); nd = b.GetNoDataValue()
    return np.where(a == nd, np.nan, a), d.GetGeoTransform()


dom, gt = raster("dom_height_66ft.tif"); cover, _ = raster("canopy_cover_66ft.tif"); slope, _ = raster("slope_plan_pct.tif")


def sample(arr, x, y):
    return float(map_coordinates(np.nan_to_num(arr), [[(y - gt[3]) / gt[5]], [(x - gt[0]) / gt[1]]], order=1, mode="nearest")[0])


uds = ogr.Open(os.path.join(WORK, "planning.gpkg")); ul = uds.GetLayerByName("units"); sp = ul.GetSpatialRef()
units = [(dict((k, f.GetField(k)) for k in ("unit_id", "method", "acres", "net_acres")), f.GetGeometryRef().Clone()) for f in ul]   # net_acres = gross less stream exclusion zones (03)

# ---- simulate plots and trees ----
plots, trees = [], []
for u, g in units:
    e = g.GetEnvelope(); x0 = math.floor(e[0] / GRID) * GRID; y0 = math.floor(e[2] / GRID) * GRID; n = 0
    for x in np.arange(x0, e[1] + GRID, GRID):
        for y in np.arange(y0, e[3] + GRID, GRID):
            if not g.Contains(ogr.CreateGeometryFromWkt(f"POINT ({x} {y})")):
                continue
            n += 1; pid = f"{u['unit_id']}-{n:02d}"
            dh = max(30.0, sample(dom, x, y)); cv = min(0.95, max(0.05, sample(cover, x, y)))
            ba_target = 60 + 180 * cv + rng.normal(0, 25)          # sq ft/ac, cover-driven
            ntrees = max(0, int(round(ba_target / BAF)))
            # ti = index range of this plot's trees in the trees list; per-plot volume is summed by index, not by plot id, so the
            # planted duplicate id cannot double-assign a plot's trees (adds no random draws: the tree list is unchanged)
            plots.append(dict(plot=pid, unit_id=u["unit_id"], x=float(x), y=float(y), n_trees=ntrees, slope=round(sample(slope, x, y)), cruiser="SIM", date="simulated", dom_ht=round(dh), cover=cv * 100, ti=(len(trees), len(trees) + ntrees)))
            for t in range(ntrees):
                sp_code = rng.choice(list(SPECIES), p=[v[1] for v in SPECIES.values()])
                dbh = float(np.clip(rng.lognormal(math.log(0.28 * dh), 0.35), 5, 60))
                ht = float(np.clip(4.5 + (dh - 4.5) * (dbh / (0.28 * dh)) ** 0.45 + rng.normal(0, 8), 10, 220))
                trees.append(dict(plot=pid, unit_id=u["unit_id"], tree=t + 1, species=sp_code, dbh=round(dbh, 1), height=round(ht), status="L", defect=int(rng.choice([0, 0, 0, 5, 10, 20], p=[0.6, 0.15, 0.1, 0.05, 0.05, 0.05]))))
print(f"{len(plots)} plots, {len(trees)} trees across {len(units)} units")

# ---- plant recording errors ----
planted = []
def plant(kind, **kw):
    planted.append(dict(kind=kind, **kw))
cand = [p for p in plots if p["n_trees"] >= 3]
p1 = cand[3]; t1 = next(t for t in trees if t["plot"] == p1["plot"]); t1["species"] = "PPP"; plant("species_code", plot=p1["plot"], tree=t1["tree"])
p2 = cand[11]; t2 = next(t for t in trees if t["plot"] == p2["plot"]); t2["dbh"] = 999.0; plant("dbh_out_of_range", plot=p2["plot"], tree=t2["tree"])
p3 = cand[19]; p3["x"] += 900; plant("plot_outside_unit", plot=p3["plot"])
p4 = cand[27]; p4["plot"] = cand[26]["plot"]; plant("duplicate_plot_id", plot=p4["plot"])
p5 = cand[35]; t5 = [t for t in trees if t["plot"] == p5["plot"]][1]; t5["height"] = None; plant("missing_height", plot=p5["plot"], tree=t5["tree"])
p6 = cand[43]; t6 = [t for t in trees if t["plot"] == p6["plot"]][0]; t6["height"] = 15; t6["dbh"] = 34.0; plant("height_dbh_mismatch", plot=p6["plot"], tree=t6["tree"])
json.dump(planted, open(os.path.join(WORK, "qc_planted.json"), "w"), indent=1)

# ---- QA checks ----
ugeom = {u["unit_id"]: g for u, g in units}
flags = {}
def flag(pid, msg):
    flags.setdefault(pid, []).append(msg)
seen = {}
for p in plots:
    seen.setdefault(p["plot"], 0); seen[p["plot"]] += 1
for pid, c in seen.items():
    if c > 1:
        flag(pid, f"Plot id recorded {c} times. Check plot numbers.")
for p in plots:
    if not ugeom[p["unit_id"]].Contains(ogr.CreateGeometryFromWkt(f"POINT ({p['x']} {p['y']})")):
        flag(p["plot"], "Plot falls outside the unit boundary. Check GPS position and unit number.")
for t in trees:
    if t["species"] not in SPECIES:
        flag(t["plot"], f"Tree {t['tree']}: species code '{t['species']}' not valid. Correct in edit.")
    if t["dbh"] < 1 or t["dbh"] > 80:
        flag(t["plot"], f"Tree {t['tree']}: DBH {t['dbh']} out of range. Remeasure.")
    if t["height"] is None:
        flag(t["plot"], f"Tree {t['tree']}: height missing. Measure.")
    elif t["dbh"] >= 20 and t["height"] < 40:
        flag(t["plot"], f"Tree {t['tree']}: height {t['height']} ft outside tolerance for DBH {t['dbh']}. Remeasure.")
found = set(flags)
print(f"QA flagged {len(found)} plots; planted errors caught: {sum(1 for e in planted if e['plot'] in found)} of {len(planted)}")

# ---- stand metrics per unit ----
# Standards and sources (all cited in docs/methods.md and on the Standards and assumptions sheet):
#   basal area per tree 0.005454 x DBH^2; BAF expansion; SDI in the summation form (Shaw 2000), maximum = basal-area-weighted
#   average of the FVS Western Sierra variant species maxima (table 3.5.1, revision of 2025-09-23), relative-density zones 35 %
#   (full occupancy) and 60 % (competition mortality) after Long (1985); total-stem cubic volume (CVTS) by the PNW-FIA species
#   equations for California (MacLean and Berger 1976, PNW-266, as tabulated in the CARB "Volume Equations" compendium), with a
#   form-factor fallback for unknown codes (Avery and Burkhart); Scribner board feet at BF_PER_CF (Keegan et al. 2010, table 2,
#   California 2000-2006); green weight by species (Miles and Smith 2009, NRS-38 table 1A); sampling error at 95 % with
#   Student's t on n-1 df computed on plot net cubic VOLUME per acre (the basal-area error is kept as a secondary figure),
#   tested against the sampling error standards a Plumas sale is designed to: Region 5 Supplement 2409.12-2021-4 to FSH
#   2409.12 ch. 40 (effective 2021-08-11), which replaces the national sec. 41.1 figures. Sec. 41.1 stratum volume error
#   standard: "The maximum sampling error for any one stratum that represents more than 10 percent of the total sale volume
#   is 40 percent for scaled sales and 30 percent for tree measurement sales" (national text 41.1(5)(b): 40 %). Sale as a
#   whole, sec. 41.1 and R5 exhibit 01: $250,000 or more, "for a scaled sale must be 20 percent or less, for tree
#   measurement sale must be 10 percent or less"; $2,000 to $5,000, 35 % scaled / 25 % tree measurement; under $2,000 with
#   sawlogs, the stratum figures; "See FSH 2409.12 Chapter 40, 41.1 Exhibit 01 for sales more than $5,000" (the national
#   tiers, up to $250,000). Sec. 41.3: "each major species or species group shall have at least 20 measured trees. A major
#   species is defined as one that comprises 10 percent or more of the sale value" and "For area based (plot) cruises,
#   there shall be a minimum of 20 plots". The sale is placed by value with the Region 5 FY2025 average sold value per MBF
#   (Cut and Sold report) on net acres. Plots tally all trees, so this is a pre-marking stand exam; the sale-cruise
#   standards are applied as if the whole tallied volume were designated for cutting.
from scipy import stats
SDI_MAX_SP = {"PP": 365, "WF": 800, "DF": 570, "SP": 561, "IC": 576}     # FVS WS variant overview table 3.5.1, revision of 2025-09-23
SDI_SRC_URL = "https://www.fs.usda.gov/sites/default/files/forest-management/fvs-ws-overview.pdf"
R5_SUPP = "R5 Supplement 2409.12-2021-4"                                    # Pacific Southwest Region supplement to FSH 2409.12 ch. 40, effective 2021-08-11
R5_SUPP_URL = "https://www.fs.usda.gov/im/directives/field/r5/fsh/2409.12/R5-2409-12-40-2021-4.docx"
STRATUM_STD = 30.0                                                          # R5 41.1 stratum VOLUME error standard, tree-measurement sales (40 % scaled)
STRATUM_STD_SCALED = 40.0                                                   # R5 41.1 stratum standard, scaled sales
STRATUM_STD_NATIONAL = 40.0                                                 # national text FSH 2409.12 41.1(5)(b), tree measurement; superseded in R5
STRATUM_SHARE_PCT = 10.0                                                    # R5 41.1: the stratum standard binds for strata over 10 % of total sale volume
SALE_TIERS = [(10000, 25), (20000, 20), (45000, 18), (70000, 16), (95000, 14), (120000, 12), (float("inf"), 10)]   # national 41.1 exhibit 01, tree measurement; R5 exhibit 01 adopts it for sales between $5,000 and $250,000
R5_HIGH_VALUE, R5_HIGH_VALUE_STD = 250000.0, 10.0                           # R5 41.1 high-value sale, tree measurement (20 % scaled)
R5_SMALL_SALE, R5_SMALL_SALE_STD = 5000.0, 25.0                             # R5 exhibit 01 small timber sale $2,000 to $5,000, tree measurement (35 % scaled)
R5_LOW_VALUE = 2000.0                                                       # R5 41.1 low-value timber sale (sawlogs) under $2,000: the stratum figure applies
MIN_MEASURED_TREES, MAJOR_SPECIES_PCT = 20, 10.0                            # R5 41.3: 20 measured trees per major species; major = 10 % or more of sale value


def sale_standard(value):
    """R5 41.1 sale-as-a-whole volume error standard (percent at 95 %) for a tree-measurement sale of the given estimated value,
    with the basis in words and the national exhibit 01 tier for comparison."""
    nat = next(pct for lim, pct in SALE_TIERS if value < lim)
    if value >= R5_HIGH_VALUE:
        return R5_HIGH_VALUE_STD, f"{R5_SUPP} sec. 41.1, high-value sale of $250,000 or more (10 % tree measurement, 20 % scaled)", nat
    if value < R5_LOW_VALUE:
        return STRATUM_STD, f"{R5_SUPP} sec. 41.1, low-value timber sale under $2,000 with sawlogs (the stratum figure)", nat
    if value < R5_SMALL_SALE:
        return R5_SMALL_SALE_STD, f"{R5_SUPP} exhibit 01, small timber sale of $2,000 to $5,000 (25 % tree measurement, 35 % scaled)", nat
    return nat, f"national FSH 2409.12 41.1 exhibit 01 tier, which {R5_SUPP} exhibit 01 adopts for sales between $5,000 and $250,000", nat
# Stumpage used to place the sale in exhibit 01: Region 5 total sold value / total sold volume, FY2025 Q1-Q4 Cut and Sold report
# (CUTS203R, run 2025-12-08). The report notes that Good Neighbor sale values are excluded, so this is a demonstration placement.
R5_FY2025_SOLD_VALUE, R5_FY2025_SOLD_MBF = 9223272.19, 274275.08
STUMPAGE_PER_MBF = R5_FY2025_SOLD_VALUE / R5_FY2025_SOLD_MBF                # $33.63/MBF
STUMPAGE_SRC_URL = "https://www.fs.usda.gov/sites/default/files/2025-q4-cut-sold-r05.pdf"
BF_PER_CF = 5.02                # Keegan et al. 2010, Forest Prod. J. 60(2):133-139, table 2 (p. 135): California, 2000-2006, Scribner bf per cu ft of bole wood
BF_SRC_URL = "https://research.fs.usda.gov/treesearch/37833"
FORM_FACTOR = 0.42              # fallback only, for species codes without a CVTS equation (conifer form factors after Avery and Burkhart, Forest Measurements)
GREEN_LB_PER_CF = {"PP": 45.0, "WF": 47.0, "DF": 38.0, "SP": 49.0, "IC": 45.0}   # Miles and Smith 2009, NRS-38 table 1A (pp. 8-9): green weight of wood, lb/cf
GREEN_LB_DEFAULT = 45.0         # fallback for unknown species codes: rounded mean of the five species above
GREEN_SRC_URL = "https://research.fs.usda.gov/treesearch/34185"
CVTS_SRC_URL = "https://ww2.arb.ca.gov/sites/default/files/cap-and-trade/protocols/usforest/2011/volume_equations.pdf"
DBH_CLASS = 4                                                               # stand table class width, in
PRACTICE_MIN_PLOTS = MIN_PLOTS                                              # R5 41.3 minimum of 20 plots for an area-based (plot) cruise


def cvts(sp, dbh, ht):
    """Total-stem cubic volume (CVTS, cu ft, including top and stump) by the PNW-FIA equations for California species:
    DF eq. 3, PP eq. 5, IC eq. 19, SP eq. 20, WF eq. 23 (MacLean and Berger 1976, PNW-266), as tabulated in the CARB
    "Volume Equations" compendium: species table p. 5 (CA column), equation pages 9, 11, 25, 26 and 29 (CVTS_SRC_URL).
    Trees under 6 in DBH use TMP_DBH = 6 and the small-tree tarif. The CF4 bounds below (0.30 to 0.40; incense-cedar floor
    0.27 with no ceiling) are printed on the CARB equation pages for each species and are applied as printed; they are not a
    guard added by this project. Returns None for species codes without an equation."""
    td = max(dbh, 6.0); ba = 0.005454154 * td ** 2
    if sp == "DF":
        cf4 = min(0.4, max(0.3, 0.248569 + 0.0253524 * ht / td - 0.0000560175 * ht ** 2 / td))
    elif sp == "PP":
        cf4 = min(0.4, max(0.3, 0.402060 - 0.899914 / td))
    elif sp == "SP":
        cf4 = min(0.4, max(0.3, 0.358550 - 0.488134 / td))
    elif sp == "WF":
        cf4 = min(0.4, max(0.3, 0.299039 + 1.91272 / ht + 0.0000367217 * ht ** 2 / td))
    elif sp == "IC":
        cf4 = max(0.27, 0.225786 + 4.44236 / ht)
    else:
        return None
    cv4 = 0.005454154 * td ** 2 * ht * cf4                              # cu ft above stump to a 4 in top
    tarif = cv4 * 0.912733 / (ba - 0.087266)
    top = 1.033 * (1.0 + 1.382937 * math.exp(-4.015292 * dbh / 10.0)) * (ba + 0.087266) - 0.174533
    if dbh > 6.0:
        return cv4 * top / (ba - 0.087266)
    small = max(0.01, 0.5 * (6.0 - dbh) ** 2 + (1.0 + 0.063 * (6.0 - dbh) ** 2) * tarif)
    return small * top


def tree_vol(t, ef):
    """Net cubic feet per acre for one tally tree: CVTS by species, form-factor approximation for unknown codes."""
    h = t["height"] or 0
    if h <= 0:
        return 0.0
    v = cvts(t["species"], t["dbh"], h)
    if v is None:
        v = 0.005454 * t["dbh"] ** 2 * h * FORM_FACTOR
    return ef * v * (1 - t["defect"] / 100)


def err_stats(vals):
    """Cruise statistics for a list of per-plot values (basal area or net cubic volume per acre): mean, SD, one SE as a
    percent of the mean, Student's t at 0.975 on n-1 df, CV percent, sampling error percent at 95 % confidence, and the
    plots needed to reach the stratum standard, (t x CV / E)^2 (FSH 2409.12 ch. 30 sample size; ch. 40 sec. 41.1)."""
    n = len(vals); mean = float(np.mean(vals)); sd = float(np.std(vals, ddof=1)) if n > 1 else 0.0; se = sd / math.sqrt(n)
    tval = float(stats.t.ppf(0.975, n - 1)) if n > 1 else 2.0
    cv = sd / mean * 100 if mean else 0.0; se95 = tval * se / mean * 100 if mean else 0.0
    return dict(mean=mean, sd=sd, se_pct=(se / mean * 100 if mean else 0.0), tval=tval, cv=cv, se95=se95, needed=(int(math.ceil((tval * cv / STRATUM_STD) ** 2)) if cv else 1))


def metrics(uid):
    ps = [p for p in plots if p["unit_id"] == uid]; n = len(ps)
    if n == 0:
        return None
    ba_plot = [p["n_trees"] * BAF for p in ps]
    # tally trees per plot by simulation index (p["ti"]); the same DBH range as the QA check, so the planted 999 in DBH drops out
    plot_trees = [[t for t in trees[p["ti"][0]:p["ti"][1]] if 1 <= t["dbh"] <= 80] for p in ps]
    # per-plot net cubic volume per acre = sum over the plot's tally trees of net CVTS x (BAF / tree BA); the unit volume below
    # is the same sum with the expansion divided by n, so mean(vol_plot) equals the unit cu ft/ac
    vol_plot = [sum(tree_vol(t, BAF / (0.005454 * t["dbh"] ** 2)) for t in pt) for pt in plot_trees]
    tl = [t for pt in plot_trees for t in pt]
    tpa = vol = sumd2 = sdi = 0.0; sp_ba = {}; sp_stat = {}; cls_stat = {}
    saw = dict(n=0, ba=0.0, tpa=0.0, d2=0.0, vol=0.0, lb=0.0); bio = dict(n=0, ba=0.0, tpa=0.0, d2=0.0, vol=0.0, lb=0.0)
    for t in tl:
        ef = BAF / (0.005454 * t["dbh"] ** 2) / n; v = tree_vol(t, ef)
        tpa += ef; sumd2 += ef * t["dbh"] ** 2; vol += v; sdi += ef * (t["dbh"] / 10.0) ** 1.605
        sp_ba[t["species"]] = sp_ba.get(t["species"], 0) + BAF / n
        g = saw if t["dbh"] >= 10 else bio; g["n"] += 1; g["ba"] += BAF / n; g["tpa"] += ef; g["d2"] += ef * t["dbh"] ** 2; g["vol"] += v; g["lb"] += v * GREEN_LB_PER_CF.get(t["species"], GREEN_LB_DEFAULT)
        s = sp_stat.setdefault(t["species"], dict(n=0, tpa=0.0, ba=0.0, d2=0.0, vol=0.0)); s["n"] += 1; s["tpa"] += ef; s["ba"] += BAF / n; s["d2"] += ef * t["dbh"] ** 2; s["vol"] += v
        k = int(t["dbh"] // DBH_CLASS) * DBH_CLASS
        c = cls_stat.setdefault(k, dict(n=0, tpa=0.0, ba=0.0, vol=0.0)); c["n"] += 1; c["tpa"] += ef; c["ba"] += BAF / n; c["vol"] += v
    for g in (saw, bio):
        g["qmd"] = math.sqrt(g["d2"] / g["tpa"]) if g["tpa"] else 0.0
    for s in sp_stat.values():
        s["qmd"] = math.sqrt(s["d2"] / s["tpa"]) if s["tpa"] else 0.0
    bae = err_stats(ba_plot); ve = err_stats(vol_plot); ba = bae["mean"]      # ve: the standards apply to the volume error (41.1(5)(a), (b))
    qmd = math.sqrt(sumd2 / tpa) if tpa else 0
    known = {k: v for k, v in sp_ba.items() if k in SDI_MAX_SP}
    sdimax = sum(v * SDI_MAX_SP[k] for k, v in known.items()) / sum(known.values()) if known else float(np.mean(list(SDI_MAX_SP.values())))
    sdi_pct = sdi / sdimax * 100 if sdimax else 0.0
    size_cls = "1" if qmd < 1 else "2" if qmd < 6 else "3" if qmd < 11 else "4" if qmd < 24 else "5"
    cover = float(np.mean([p["cover"] for p in ps]))
    dens_cls = "open" if cover < 10 else "S" if cover < 25 else "P" if cover < 40 else "M" if cover < 60 else "D"   # CWHR: S 10-24, P 25-39, M 40-59, D 60-100
    # measured trees by species for the R5 41.3 check: every simulated tally tree is measured for DBH and height, so the count is
    # the tally less the records QA drops (bad species code, DBH out of range, missing height)
    measured = {}
    for t in tl:
        if t["species"] in SPECIES and t["height"] is not None:
            measured[t["species"]] = measured.get(t["species"], 0) + 1
    return dict(plots=n, ba=ba, tval=ve["tval"], se_std=STRATUM_STD, measured=measured,
                ba_sd=bae["sd"], ba_se_pct=bae["se_pct"], ba_cv=bae["cv"], ba_se95=bae["se95"], ba_needed=bae["needed"], ba_meets=(bae["se95"] <= STRATUM_STD),
                vol_sd=ve["sd"], vol_se_pct=ve["se_pct"], vol_cv=ve["cv"], vol_se95=ve["se95"], vol_needed=ve["needed"],
                meets=(ve["se95"] <= STRATUM_STD), practice_ok=(n >= PRACTICE_MIN_PLOTS), tpa=tpa, qmd=qmd, cuft=vol, mbf=vol / 1000 * BF_PER_CF,
                species=sorted(sp_ba.items(), key=lambda x: -x[1]), sp_stat=sp_stat, cls_stat=cls_stat, sdi=sdi, sdimax=sdimax, sdi_pct=sdi_pct,
                size_cls=size_cls, dens_cls=dens_cls, cover=cover, saw=saw, bio=bio, target_ba=ba * TARGET_SDI_PCT / max(1e-6, sdi_pct))


M = {u["unit_id"]: metrics(u["unit_id"]) for u, g in units}
# sale as a whole: stratified by unit with gross-area weights, t = 2 (FSH 2409.12 41.1(5)(a): the sale-as-a-whole standard is a
# volume error; the basal-area figure is kept alongside). Plots sample the gross unit, so per-acre means are weighted by gross
# acres; sale MBF and value for the exhibit 01 placement use NET acres (gross less stream exclusion zones), gross also reported.
tot_ac = sum(u["acres"] for u, g in units); tot_net = int(round(sum(round(u["net_acres"], 2) for u, g in units), 2) + 0.5)   # half-up, matches the overview sheet
est = [(u["acres"] / tot_ac, M[u["unit_id"]], u) for u, g in units if M[u["unit_id"]]]
sale_ba = sum(w * m["ba"] for w, m, u in est); sale_ba_se95 = 2 * math.sqrt(sum(w ** 2 * m["ba_sd"] ** 2 / m["plots"] for w, m, u in est)) / sale_ba * 100
sale_vol = sum(w * m["cuft"] for w, m, u in est); sale_vol_se95 = 2 * math.sqrt(sum(w ** 2 * m["vol_sd"] ** 2 / m["plots"] for w, m, u in est)) / sale_vol * 100
sale_mbf_gross = sum(u["acres"] * m["mbf"] for w, m, u in est); sale_value_gross = sale_mbf_gross * STUMPAGE_PER_MBF
sale_mbf = sum(u["net_acres"] * m["mbf"] for w, m, u in est); sale_value = sale_mbf * STUMPAGE_PER_MBF
sale_std, sale_basis, sale_std_national = sale_standard(sale_value); sale_meets = sale_vol_se95 <= sale_std
print(f"sale as a whole: volume {sale_vol:,.0f} cu ft/ac, volume SE {sale_vol_se95:.1f} % at 95 % (BA {sale_ba:.0f} sq ft/ac, BA SE {sale_ba_se95:.1f} %), "
      f"standard {sale_std} % ({'MEETS' if sale_meets else 'FAILS'}; {sale_basis}; national exhibit 01 tier {sale_std_national} %) for a net-acre value of ${sale_value:,.0f} ({sale_mbf:,.0f} MBF on {tot_net:,.0f} net ac; "
      f"gross-acre total {sale_mbf_gross:,.0f} MBF on {tot_ac:,.0f} ac = ${sale_value_gross:,.0f}) at the Region 5 FY2025 sold average of ${STUMPAGE_PER_MBF:.2f}/MBF")
# R5 41.1: the stratum standard binds for any stratum over STRATUM_SHARE_PCT of total sale volume; every unit is tested, and the
# sheet says whether the standard binds for it or is shown for information. Share = the unit's net-acre MBF over the sale total.
for w, m, u in est:
    m["share"] = u["net_acres"] * m["mbf"] / sale_mbf * 100; m["binds"] = m["share"] > STRATUM_SHARE_PCT
# R5 41.3: at least 20 measured trees per major species; major = 10 % or more of sale VALUE. One stumpage rate is applied to
# every species, so value share equals net-acre volume share. Counted for the sale (the requirement) and for each unit.
sale_sp_vol = {}
for w, m, u in est:
    for k, s in m["sp_stat"].items():
        sale_sp_vol[k] = sale_sp_vol.get(k, 0.0) + u["net_acres"] * s["vol"]
sale_sp_share = {k: v / sum(sale_sp_vol.values()) * 100 for k, v in sale_sp_vol.items()}
major_species = [k for k, v in sorted(sale_sp_share.items(), key=lambda x: -x[1]) if k in SPECIES and v >= MAJOR_SPECIES_PCT]
sale_measured = {k: sum(m["measured"].get(k, 0) for w, m, u in est) for k in SPECIES}
sale_trees_ok = all(sale_measured[k] >= MIN_MEASURED_TREES for k in major_species)
for w, m, u in est:
    m["trees_short"] = [k for k in major_species if m["measured"].get(k, 0) < MIN_MEASURED_TREES]; m["trees_ok"] = not m["trees_short"]
print(f"major species (>= {MAJOR_SPECIES_PCT:.0f} % of sale value): " + ", ".join(f"{k} {sale_sp_share[k]:.1f} %" for k in major_species)
      + f"; measured trees for the sale: " + ", ".join(f"{k} {sale_measured[k]}" for k in SPECIES) + f" ({'MEETS' if sale_trees_ok else 'FAILS'} the R5 41.3 minimum of {MIN_MEASURED_TREES}); "
      f"units meeting it: {sum(1 for w, m, u in est if m['trees_ok'])} of {len(est)}; units where the stratum standard binds (> {STRATUM_SHARE_PCT:.0f} % of sale volume): {sum(1 for w, m, u in est if m['binds'])}; "
      f"units meeting {STRATUM_STD:.0f} %: {sum(1 for w, m, u in est if m['meets'])} of {len(est)}, worst {max(m['vol_se95'] for w, m, u in est):.1f} %; units with 20 plots: {sum(1 for w, m, u in est if m['practice_ok'])}")

# ---- outputs: gpkg, xlsx ----
drv = ogr.GetDriverByName("GPKG"); gp = os.path.join(WORK, "cruise_plots.gpkg")
if os.path.exists(gp):
    drv.DeleteDataSource(gp)
ds = drv.CreateDataSource(gp); pl = ds.CreateLayer("plots", sp, ogr.wkbPoint)
for n, t in (("plot", ogr.OFTString), ("unit_id", ogr.OFTInteger), ("n_trees", ogr.OFTInteger), ("slope", ogr.OFTInteger), ("flags", ogr.OFTString)):
    pl.CreateField(ogr.FieldDefn(n, t))
for p in plots:
    f = ogr.Feature(pl.GetLayerDefn()); f.SetGeometry(ogr.CreateGeometryFromWkt(f"POINT ({p['x']} {p['y']})"))
    f.SetField("plot", p["plot"]); f.SetField("unit_id", p["unit_id"]); f.SetField("n_trees", p["n_trees"]); f.SetField("slope", p["slope"]); f.SetField("flags", "; ".join(flags.get(p["plot"], []))); pl.CreateFeature(f)
ds = None

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
wb = Workbook(); ws = wb.active; ws.title = "Plots"; ws.append(["plot", "unit_id", "x_ft", "y_ft", "trees_in", "BA_sqft_ac", "slope_pct", "cruiser", "date", "qa_flags"])
for p in plots:
    ws.append([p["plot"], p["unit_id"], round(p["x"]), round(p["y"]), p["n_trees"], p["n_trees"] * BAF, p["slope"], p["cruiser"], p["date"], "; ".join(flags.get(p["plot"], []))])
ws2 = wb.create_sheet("Trees"); ws2.append(["plot", "unit_id", "tree", "species", "dbh_in", "height_ft", "status", "defect_pct", "BA_sqft", "expansion_TPA_per_plot"])
for t in trees:
    ws2.append([t["plot"], t["unit_id"], t["tree"], t["species"], t["dbh"], t["height"], t["status"], t["defect"], round(0.005454 * t["dbh"] ** 2, 3) if 1 <= t["dbh"] <= 80 else None, round(BAF / (0.005454 * t["dbh"] ** 2), 2) if 1 <= t["dbh"] <= 80 else None])
ws3 = wb.create_sheet("Unit summary")
ws3.append(["unit_id", "method", "gross acres", "net acres", "plots", "t (95 %, n-1)", "cu ft/ac", "SD of plot cu ft/ac", "vol_cv %", "vol_se95 %", "stratum standard % (volume, R5 41.1 tree measurement)", "meets standard (volume)", "plots for standard (volume)",
            "share of sale volume % (net acres)", "stratum standard binds (share > 10 %, R5 sec. 41.1)", "measured trees by species", "R5 20 measured trees per major species (sec. 41.3)",
            "BA sq ft/ac", "SD of plot BA", "ba_cv %", "ba_se95 %", "plots for standard (BA)", "R5 20-plot minimum (sec. 41.3)",
            "TPA", "QMD in", "MBF/ac Scribner", "SDI", "SDI max (BA-wtd FVS WS)", "SDI % of max", "CWHR size", "CWHR density", "Sawtimber BA", "Biomass BA", "Biomass green tons/ac", "QA flags"])
for u, g in units:
    m = M[u["unit_id"]]
    if m:
        nf = len({(p["plot"], f) for p in plots if p["unit_id"] == u["unit_id"] for f in flags.get(p["plot"], [])})
        ws3.append([u["unit_id"], u["method"], round(u["acres"], 1), round(u["net_acres"], 1), m["plots"], round(m["tval"], 3), round(m["cuft"]), round(m["vol_sd"]), round(m["vol_cv"], 1), round(m["vol_se95"], 1), m["se_std"], "yes" if m["meets"] else "no", m["vol_needed"],
                    round(m["share"], 1), "binds" if m["binds"] else "for information", ", ".join(f"{k} {m['measured'].get(k, 0)}" for k in SPECIES), "yes" if m["trees_ok"] else "no: " + ", ".join(m["trees_short"]),
                    round(m["ba"], 1), round(m["ba_sd"], 1), round(m["ba_cv"], 1), round(m["ba_se95"], 1), m["ba_needed"], "yes" if m["practice_ok"] else "no",
                    round(m["tpa"]), round(m["qmd"], 1), round(m["mbf"], 1), round(m["sdi"]), round(m["sdimax"]), round(m["sdi_pct"]), m["size_cls"], m["dens_cls"], round(m["saw"]["ba"]), round(m["bio"]["ba"]), round(m["bio"]["lb"] / 2000, 1), nf])
ws4 = wb.create_sheet("Stand tables"); ws4.append(["unit_id", "DBH class (in)", "trees tallied", "TPA", "BA sq ft/ac", "cu ft/ac", "MBF/ac"])
for u, g in units:
    m = M[u["unit_id"]]
    if not m:
        continue
    for k in sorted(m["cls_stat"]):
        c = m["cls_stat"][k]; ws4.append([u["unit_id"], f"{k}-{k + DBH_CLASS - 0.1:.1f}", c["n"], round(c["tpa"], 1), round(c["ba"], 1), round(c["vol"]), round(c["vol"] / 1000 * BF_PER_CF, 2)])
ws5 = wb.create_sheet("Stock tables"); ws5.append(["unit_id", "species", "trees tallied", "TPA", "BA sq ft/ac", "QMD in", "cu ft/ac", "MBF/ac"])
for u, g in units:
    m = M[u["unit_id"]]
    if not m:
        continue
    for k, s in sorted(m["sp_stat"].items(), key=lambda x: -x[1]["ba"]):
        ws5.append([u["unit_id"], k, s["n"], round(s["tpa"], 1), round(s["ba"], 1), round(s["qmd"], 1), round(s["vol"]), round(s["vol"] / 1000 * BF_PER_CF, 2)])
ws6 = wb.create_sheet("Standards and assumptions")
FSH_URL = "https://www.fs.usda.gov/im/directives/dughtml/fsh.html"   # Forest Service Handbook directives index (FSH 2409.12 ch. 40, WO amendment 2012-1)
for row in (["item", "value", "source", "url"],
            ["Basal area factor", BAF, "variable-radius (prism) cruise; BA/ac = trees in x BAF (Univ. of Tennessee Extension W1117, common forest measurements)", "https://utia.tennessee.edu/publications/wp-content/uploads/sites/269/2023/10/W1117.pdf"],
            ["Tree basal area", "0.005454 x DBH^2 sq ft", "standard mensuration (Univ. of Tennessee Extension W1117)", "https://utia.tennessee.edu/publications/wp-content/uploads/sites/269/2023/10/W1117.pdf"],
            ["Expansion factor", "BAF / tree BA, divided by plots", "per-tree trees per acre in a variable-radius cruise (Univ. of Tennessee Extension W1117)", "https://utia.tennessee.edu/publications/wp-content/uploads/sites/269/2023/10/W1117.pdf"],
            ["Sampling error (volume)", "t(0.975, n-1) x SE / mean of plot net cubic volume per acre, percent; plot volume = sum over tally trees of net CVTS x BAF / tree BA (sale as a whole uses t = 2, the handbook's large-sample value)", f"{R5_SUPP} sec. 41.1 (effective 2021-08-11), which replaces the national FSH 2409.12 ch. 40 sec. 41.1(5) figures for a Region 5 sale: the sale-as-a-whole and stratum error standards apply to VOLUME at 95 % confidence (t = 2 for large n)", R5_SUPP_URL],
            ["Sampling error (basal area)", "the same statistic on plot basal area (trees in x BAF); reported as a secondary figure, not tested against the standard", "retained for reference; volume CV runs about twice the basal-area CV, so the basal-area error understates what the handbook tests", R5_SUPP_URL],
            ["Standards source", f"{R5_SUPP}, FSH 2409.12 ch. 40, Pacific Southwest Region, effective 2021-08-11 (supersedes 2409.12-2019-1)", "Region 5 supplement to the Timber Cruising Handbook; national text FSH 2409.12 ch. 40 WO amendment 2012-1 at the directives index", R5_SUPP_URL + " ; " + FSH_URL],
            ["Stratum standard", f"{STRATUM_STD:.0f} % of volume, tree-measurement sales ({STRATUM_STD_SCALED:.0f} % scaled); binds for any stratum over {STRATUM_SHARE_PCT:.0f} % of total sale volume, shown for information for smaller units; national text 41.1(5)(b): {STRATUM_STD_NATIONAL:.0f} %", f"{R5_SUPP} sec. 41.1, Stratum Volume Error Standard: 'The maximum sampling error for any one stratum that represents more than 10 percent of the total sale volume is 40 percent for scaled sales and 30 percent for tree measurement sales'", R5_SUPP_URL],
            ["Sale-as-a-whole standard", f"{sale_std} % of volume at a net-acre value of ${sale_value:,.0f} ({sale_mbf:,.0f} MBF on {tot_net:,.0f} net ac x ${STUMPAGE_PER_MBF:.2f}/MBF; gross-acre total {sale_mbf_gross:,.0f} MBF on {tot_ac:,.0f} ac = ${sale_value_gross:,.0f}); national exhibit 01 tier {sale_std_national} %", f"{sale_basis}. R5 41.1 and exhibit 01: $250,000 or more, 'for a scaled sale must be 20 percent or less, for tree measurement sale must be 10 percent or less'; $2,000 to $5,000, 35 % scaled / 25 % tree measurement; under $2,000 with sawlogs, the stratum figures; 'See FSH 2409.12 Chapter 40, 41.1 Exhibit 01 for sales more than $5,000'. Net acres = gross less stream equipment exclusion zones", R5_SUPP_URL],
            ["Stumpage for sale value", f"${STUMPAGE_PER_MBF:.2f} per MBF = Region 5 sold value ${R5_FY2025_SOLD_VALUE:,.2f} / sold volume {R5_FY2025_SOLD_MBF:,.2f} MBF, all sales", "Forest Service Cut and Sold report CUTS203R, Region 5, cumulative FY2025 Q1-Q4 (run 2025-12-08), region total row; the report excludes Good Neighbor sale values and says not to use it for unit values, so this is a demonstration placement", STUMPAGE_SRC_URL],
            ["Sale-as-a-whole estimate", "stratified by unit, gross-area weights (plots sample the gross unit); var = sum(W^2 s^2 / n) on plot volume per acre, basal area alongside", "Cochran; FSH 2409.12 ch. 30", FSH_URL],
            ["R5 20 measured trees per major species", f"{MIN_MEASURED_TREES} measured trees for each species with {MAJOR_SPECIES_PCT:.0f} % or more of sale value; major species here: " + ", ".join(f"{k} {sale_sp_share[k]:.1f} %" for k in major_species) + "; sale counts: " + ", ".join(f"{k} {sale_measured[k]}" for k in SPECIES), f"{R5_SUPP} sec. 41.3: 'each major species or species group shall have at least 20 measured trees. A major species is defined as one that comprises 10 percent or more of the sale value'. One stumpage rate is applied to every species, so value share equals net-acre volume share. Every simulated tally tree is measured for DBH and height; QA-dropped records are not counted", R5_SUPP_URL],
            ["R5 20-plot minimum", PRACTICE_MIN_PLOTS, f"{R5_SUPP} sec. 41.3: 'For area based (plot) cruises, there shall be a minimum of 20 plots'", R5_SUPP_URL],
            ["Cruise type", "pre-marking stand exam: plots tally all trees; the sale-cruise standards are applied as if the whole tallied volume were designated for cutting", "demonstration convention; a sale cruise would tally designated trees only", "https://github.com/woodsy-will/plumas-lidar-harvest-planning/blob/main/docs/methods.md"],
            ["Plots for standard", f"(t x CV / E)^2 with the volume CV and E = {STRATUM_STD:.0f} % (the basal-area figure is also listed)", "FSH 2409.12 ch. 30 sample size", FSH_URL],
            ["SDI", "sum over trees of TPA x (DBH/10)^1.605", "Reineke 1933; summation form after Shaw 2000, reviewed in Shaw 2006", "https://www.fs.usda.gov/rm/pubs_other/rmrs_2006_shaw_j006.pdf"],
            ["SDI maximum", "BA-weighted mean of species maxima: " + ", ".join(f"{k} {v}" for k, v in SDI_MAX_SP.items()), "FVS Staff 2008 (revised 2025-09-23), Western Sierra Nevada (WS) Variant Overview, Forest Vegetation Simulator, table 3.5.1 (sources Shaw and PSW); earlier revisions of the overview listed different maxima for some species", SDI_SRC_URL],
            ["Relative density zones", "35 % of max = lower limit of full site occupancy; 60 % = onset of competition mortality", "Long 1985; Long and Shaw 2012 (Sierra mixed conifer DMD)", "https://research.fs.usda.gov/treesearch/25003"],
            ["Leave target", f"{TARGET_SDI_PCT} % of max SDI, expressed as BA", "demonstration target at the full-occupancy threshold", "https://research.fs.usda.gov/treesearch/25003"],
            ["Cubic volume", "total-stem cubic volume (CVTS) by species: DF eq. 3, PP eq. 5, IC eq. 19, SP eq. 20, WF eq. 23 (tarif system, small-tree tarif under 6 in), net of recorded defect", "PNW-FIA volume equations for California (MacLean and Berger 1976, PNW Research Note PNW-266) as tabulated in 'Volume estimation for the PNW-FIA Integrated Database', reproduced by the California Air Resources Board (2011): species table p. 5 (CA column), equations pp. 9, 11, 25, 26, 29", CVTS_SRC_URL],
            ["CVTS form-class bounds", "CF4 held to 0.30 to 0.40 for DF, PP, SP and WF; incense-cedar floor 0.27, no ceiling", "these bounds are printed on the CARB 2011 equation pages for each species and are applied as printed; they are not a guard added by this project", CVTS_SRC_URL],
            ["Cubic volume fallback", f"BA x total height x form factor {FORM_FACTOR} for species codes without an equation (e.g. the planted bad code)", "conifer form-factor approximation after Avery and Burkhart, Forest Measurements (5th ed., McGraw-Hill 2002), tree form and volume chapters", "n/a: Avery and Burkhart, Forest Measurements (textbook)"],
            ["Board feet", f"{BF_PER_CF} Scribner bf per cu ft, applied to CVTS", "Keegan, Morgan, Blatner and Daniels 2010, Trends in lumber processing in the western United States, Part I, Forest Prod. J. 60(2):133-139, table 2 (p. 135): California, 2000-2006, board feet Scribner per cubic foot of bole wood inside bark (the abstract rounds it to 5.03). The Keegan ratio is Scribner log scale per cubic foot of delivered sawlog fiber; applying it to total-stem CVTS (top, stump and non-sawlog stems included) overstates sawlog board-foot volume. The consistent pairing would be merchantable cubic volume, which these tarif equations do not produce without a merchantable-top conversion", BF_SRC_URL],
            ["Biomass green tons", "green weight of wood by species: " + ", ".join(f"{k} {v:.0f}" for k, v in GREEN_LB_PER_CF.items()) + f" lb per cu ft; unknown codes {GREEN_LB_DEFAULT:.0f}", "Miles and Smith 2009, Specific gravity and other properties of wood and bark for 156 tree species found in North America, Research Note NRS-38, table 1A (pp. 8-9), 'Avg. green wt. (lb/cf)' of wood on a green-volume basis; Douglas-fir is a single entry there (FIA code 202, not split coast/interior); bark excluded", GREEN_SRC_URL],
            ["CWHR size class", "QMD: 1 <1 in, 2 1-6, 3 6-11, 4 11-24, 5 >24", "California Wildlife Habitat Relationships", "https://wildlife.ca.gov/Data/CWHR"],
            ["CWHR density class", "canopy cover: S 10-24 %, P 25-39, M 40-59, D 60-100", "California Wildlife Habitat Relationships", "https://wildlife.ca.gov/Data/CWHR"],
            ["Data", "SIMULATED: every plot and tree is generated from LiDAR canopy metrics with six planted recording errors; cruiser code SIM, date field 'simulated', SIMULATED DATA watermark on every review-sheet page", "docs/methods.md", "https://github.com/woodsy-will/plumas-lidar-harvest-planning/blob/main/docs/methods.md"]):
    ws6.append(row)
ws0 = wb.create_sheet("READ ME", 0)
ws0.append(["SIMULATED DATA. Every plot and tree in this workbook is simulated from LiDAR canopy metrics (dominant height and canopy cover) with six planted recording errors. "
            "It is not field data: no crew measured these trees. The cruiser code SIM and the date field 'simulated' mark every record, and the review-sheet PDFs carry a SIMULATED DATA watermark."])
ws0.append([]); ws0.append(["Sheet", "Contents"])
for nm, what in (("Plots", "one row per simulated plot: location, trees in, BA, slope, cruiser SIM, date 'simulated', QA flags"),
                 ("Trees", "one row per simulated tally tree: species, DBH, height, defect, tree BA, expansion factor"),
                 ("Unit summary", "per-unit cruise statistics, standards check, SDI, CWHR, sawtimber and biomass"),
                 ("Stand tables", "TPA, BA and volume by 4-inch DBH class per unit"),
                 ("Stock tables", "TPA, BA, QMD and volume by species per unit"),
                 ("Standards and assumptions", "every constant with its value, source and URL")):
    ws0.append([nm, what])
ws0["A1"].font = Font(bold=True); ws0["A1"].alignment = Alignment(wrap_text=True, vertical="top"); ws0.merge_cells("A1:B1"); ws0.row_dimensions[1].height = 75
ws0["A3"].font = Font(bold=True); ws0["B3"].font = Font(bold=True); ws0.column_dimensions["A"].width = 28; ws0.column_dimensions["B"].width = 100
for w in (ws, ws2, ws3, ws4, ws5, ws6):
    for c in w[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDE8D0"); c.alignment = Alignment(wrap_text=True, vertical="top")
    w.freeze_panes = "A2"
    for i, col in enumerate(w.columns, 1):
        w.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]) + 2))
wb.save(os.path.join(OUT, "cruise_data.xlsx"))

# ---- review sheets (one per unit) and the merged PDF ----
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
from reportlab.platypus import KeepTogether
FONT = "Helvetica"; FONT_B = "Helvetica-Bold"
_fd = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
if all(os.path.exists(os.path.join(_fd, f)) for f in ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")):
    for nme, f in (("Arial", "arial.ttf"), ("Arial-Bold", "arialbd.ttf"), ("Arial-Italic", "ariali.ttf"), ("Arial-BoldItalic", "arialbi.ttf")):
        pdfmetrics.registerFont(TTFont(nme, os.path.join(_fd, f)))
    addMapping("Arial", 0, 0, "Arial"); addMapping("Arial", 1, 0, "Arial-Bold"); addMapping("Arial", 0, 1, "Arial-Italic"); addMapping("Arial", 1, 1, "Arial-BoldItalic")
    FONT, FONT_B = "Arial", "Arial-Bold"
styles = getSampleStyleSheet(); H = styles["Heading2"]; B = styles["BodyText"]; B.fontSize = 9; B.leading = 11
B.fontName = FONT; H.fontName = FONT_B
from reportlab.lib.styles import ParagraphStyle
TT = ParagraphStyle("tt", parent=B, fontName=FONT_B, fontSize=8.5, leading=10, spaceBefore=6, spaceAfter=2)      # table title, above the table
FN = ParagraphStyle("fn", parent=B, fontName=FONT, fontSize=7, leading=8.5, textColor=colors.HexColor("#444444"), spaceBefore=2)   # footnote, below
CS = ParagraphStyle("cell", parent=B, fontName=FONT, fontSize=8, leading=10)                                                        # Table 1 value cell: wraps inside its column
GREEN, PALE = colors.HexColor("#dde8d0"), colors.HexColor("#f3f6ee")


def tstyle(head_rows=(0,), size=8, extra=()):
    st = [("FONTNAME", (0, 0), (-1, -1), FONT), ("FONTSIZE", (0, 0), (-1, -1), size), ("LEADING", (0, 0), (-1, -1), size + 2),
          ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black), ("LINEABOVE", (0, 0), (-1, 0), 0.6, colors.black), ("LINEBELOW", (0, -1), (-1, -1), 0.6, colors.black),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
    for r in head_rows:
        st += [("FONTNAME", (0, r), (-1, r), FONT_B), ("BACKGROUND", (0, r), (-1, r), GREEN)]
    return TableStyle(st + list(extra))


merged = SimpleDocTemplate(os.path.join(OUT, "Unit_Reviews.pdf"), pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title="Field data review sheets", author="William Steinley")
story = []


def watermark(canv, doc):
    """Light grey diagonal SIMULATED DATA on every page; the onPage callback runs before the page content is drawn."""
    w, h = doc.pagesize; canv.saveState(); canv.setFont(FONT_B, 64); canv.setFillColor(colors.HexColor("#cfcfcf"))
    canv.translate(w / 2, h / 2); canv.rotate(40); canv.drawCentredString(0, 0, "SIMULATED DATA"); canv.restoreState()


def build_page(u, g):
    uid = u["unit_id"]; m = M[uid]
    if not m:
        return None
    ps = [p for p in plots if p["unit_id"] == uid]
    fig, ax = plt.subplots(figsize=(4.2, 3.2))
    rings = [g.GetGeometryRef(i).GetGeometryRef(0) for i in range(g.GetGeometryCount())] if g.GetGeometryName() == "MULTIPOLYGON" else [g.GetGeometryRef(0)]
    for ring in rings:
        ax.plot([ring.GetPoint_2D(i)[0] for i in range(ring.GetPointCount())], [ring.GetPoint_2D(i)[1] for i in range(ring.GetPointCount())], color="k", lw=1.2)
    for p in ps:
        bad = p["plot"] in flags; ax.plot(p["x"], p["y"], "o", ms=5, color=("#D55E00" if bad else "#0072B2"), mec="k"); ax.annotate(p["plot"].split("-")[1], (p["x"], p["y"]), fontsize=5, xytext=(3, 3), textcoords="offset points")
    ax.set_aspect("equal"); ax.set_xticks([]); ax.set_yticks([]); ax.set_title(f"Unit {uid} plots, {GRID:.0f} ft grid (orange: QA flag)", fontsize=8)
    for sde in ("top", "right", "left", "bottom"):
        ax.spines[sde].set_visible(True)
    mp = os.path.join(WORK, f"_plotmap_{uid}.png"); fig.tight_layout(); fig.savefig(mp, dpi=200); plt.close(fig)
    page = [Paragraph(f"Unit {uid} Field Data Review", H),
            Paragraph(f"Mohawk Valley West Slope demonstration. <b>{u['method']}</b> unit, {u['acres']:.1f} ac. {m['plots']} plots, BAF {BAF:.0f}, variable-radius cruise on a {GRID:.0f} ft grid, {sum(p['n_trees'] for p in ps)} trees tallied. Simulated data with planted recording errors; see docs/methods.md. "
                      "Plots tally all trees, so this is a pre-marking stand exam; the sale-cruise standards are applied as if the whole tallied volume were designated for cutting.", B)]
    # Table 1: stand summary. Value cells are Paragraphs so long entries wrap inside their column instead of running into the next
    P = lambda s: Paragraph(s, CS)
    st = [["STAND DENSITY", "", "STAND INFO", ""],
          ["Basal area", P(f"{m['ba']:.0f} sq ft/ac  (1 SE of mean, BA {m['ba_se_pct']:.1f} %)"), "Type", P("Sierran mixed conifer")],
          ["Trees per acre", P(f"{m['tpa']:.0f}"), "CWHR size / density", P(f"{m['size_cls']} / {m['dens_cls']}   (cover {m['cover']:.0f} %)")],
          ["QMD", P(f"{m['qmd']:.1f} in"), "SDI", P(f"{m['sdi']:.0f}  ({m['sdi_pct']:.0f} % of max {m['sdimax']:.0f})")],
          ["SAWTIMBER  (>= 10 in)", "", "BIOMASS  (< 10 in)", ""],
          ["BA | QMD", P(f"{m['saw']['ba']:.0f} sq ft/ac | {m['saw']['qmd']:.1f} in"), "BA | QMD", P(f"{m['bio']['ba']:.0f} sq ft/ac | {m['bio']['qmd']:.1f} in")],
          ["TPA", P(f"{m['saw']['tpa']:.0f}"), "TPA", P(f"{m['bio']['tpa']:.0f}")],
          ["Volume", P(f"{m['saw']['vol'] / 1000 * BF_PER_CF:.1f} MBF/ac  ({m['saw']['vol']:.0f} cu ft)"), "Volume", P(f"{m['bio']['lb'] / 2000:.1f} green tons/ac  ({m['bio']['vol']:.0f} cu ft)")],
          ["TREATMENT TARGET", "", "CRUISE DESIGN", ""],
          ["Leave", P(f"{TARGET_SDI_PCT} % of max SDI, about {m['target_ba']:.0f} sq ft/ac BA"), "Vol. sampling error", P(f"{m['vol_se95']:.1f} % at 95 %  (CV {m['vol_cv']:.0f} %, t = {m['tval']:.2f}, n = {m['plots']})")],
          ["Species (BA)", P(", ".join(f"{s} {b:.0f}" for s, b in m["species"][:5])), "Stratum standard", P(f"{m['se_std']:.0f} % of volume (R5 2409.12-2021-4, sec. 41.1): <b>{'MEETS' if m['meets'] else 'FAILS'}</b>; plots for standard {m['vol_needed']}")],
          ["", "", "Sale volume share", P(f"{m['share']:.1f} % of sale volume: standard {'<b>binds</b> (over' if m['binds'] else 'shown for information (under'} the R5 {STRATUM_SHARE_PCT:.0f} % threshold)")],
          ["", "", "BA sampling error", P(f"{m['ba_se95']:.1f} % at 95 %  (CV {m['ba_cv']:.0f} %); reference, not the standard")],
          ["", "", "Measured trees", P(", ".join(f"{k} {m['measured'].get(k, 0)}" for k in SPECIES) + f"; R5 {MIN_MEASURED_TREES} measured trees per major species (sec. 41.3; major by sale value: {', '.join(major_species)}): " + ("<b>ok</b>" if m["trees_ok"] else "<b>SHORT</b>: " + ", ".join(f"{k} {m['measured'].get(k, 0)}" for k in m["trees_short"])))],
          ["", "", "Plot count", P(f"{m['plots']}; R5 {PRACTICE_MIN_PLOTS}-plot minimum for an area-based cruise (sec. 41.3): " + ("<b>ok</b>" if m["practice_ok"] else f"<b>SHORT</b> by {PRACTICE_MIN_PLOTS - m['plots']}"))]]
    t1 = Table(st, colWidths=[1.0 * inch, 1.9 * inch, 1.3 * inch, 3.0 * inch])
    t1.setStyle(tstyle(head_rows=(0, 4, 8), extra=[("BACKGROUND", (0, r0), (0, r1), PALE) for r0, r1 in ((1, 3), (5, 7), (9, 14))] + [("BACKGROUND", (2, r0), (2, r1), PALE) for r0, r1 in ((1, 3), (5, 7), (9, 14))]))
    page += [Paragraph("TABLE 1. STAND SUMMARY, BAF 20 CRUISE, PER ACRE", TT), t1,
             Paragraph(f"Cubic volume is total-stem CVTS by species (PNW-FIA California equations, MacLean and Berger 1976, CARB 2011 compendium), net of defect. Scribner board feet at {BF_PER_CF} bf/cu ft (Keegan et al. 2010, table 2, California 2000-2006). "
                       f"Green tons at " + ", ".join(f"{k} {v:.0f}" for k, v in GREEN_LB_PER_CF.items()) + " lb/cu ft of wood (Miles and Smith 2009, NRS-38 table 1A). SDI in the summation form; maximum is the basal-area-weighted FVS Western Sierra species maximum (table 3.5.1, rev. 2025-09-23). "
                       f"The Keegan ratio is Scribner log scale per cubic foot of delivered sawlog fiber; applied to total-stem CVTS it overstates sawlog volume. Merchantable cubic volume is the consistent pairing. "
                       f"Sampling error at 95 % confidence, Student's t, on plot net cubic volume per acre, the quantity the standards test. Standards: Region 5 Supplement 2409.12-2021-4 to FSH 2409.12 ch. 40 (effective 2021-08-11), sec. 41.1, stratum standard {STRATUM_STD:.0f} % of volume for tree-measurement sales ({STRATUM_STD_SCALED:.0f} % scaled; the national text 41.1(5)(b) reads {STRATUM_STD_NATIONAL:.0f} %), binding for any stratum over {STRATUM_SHARE_PCT:.0f} % of total sale volume and shown for information otherwise; "
                       f"sec. 41.3, {MIN_MEASURED_TREES} measured trees per major species (10 % or more of sale value; one stumpage rate, so value share equals volume share) and {PRACTICE_MIN_PLOTS} plots minimum. Measured trees are tally trees with a valid code, DBH in range and a height. The basal-area error is reference only. Stand-summary BA counts every tallied tree (plot count x BAF); the stand and stock tables drop records flagged by QA, such as an out-of-range DBH, so the two can differ by a few tenths. Sources and URLs: cruise_data.xlsx, Standards and assumptions.", FN)]
    # Table 2: stand table by DBH class; Table 3: stock table by species, side by side
    rows2 = [["DBH class, in", "Trees", "TPA", "BA", "cu ft", "MBF"]] + [[f"{k}-{k + DBH_CLASS - 1:.0f}.9", c["n"], f"{c['tpa']:.1f}", f"{c['ba']:.1f}", f"{c['vol']:.0f}", f"{c['vol'] / 1000 * BF_PER_CF:.1f}"] for k, c in sorted(m["cls_stat"].items())]
    rows2.append(["All", sum(c["n"] for c in m["cls_stat"].values()), f"{m['tpa']:.1f}", f"{m['ba']:.1f}", f"{m['cuft']:.0f}", f"{m['mbf']:.1f}"])
    t2 = Table(rows2, colWidths=[0.95 * inch, 0.5 * inch, 0.5 * inch, 0.5 * inch, 0.55 * inch, 0.5 * inch]); t2.setStyle(tstyle(extra=[("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONTNAME", (0, -1), (-1, -1), FONT_B), ("LINEABOVE", (0, -1), (-1, -1), 0.4, colors.black)]))
    rows3 = [["Species", "Trees", "TPA", "BA", "QMD", "cu ft", "MBF"]] + [[k, s["n"], f"{s['tpa']:.1f}", f"{s['ba']:.1f}", f"{s['qmd']:.1f}", f"{s['vol']:.0f}", f"{s['vol'] / 1000 * BF_PER_CF:.1f}"] for k, s in sorted(m["sp_stat"].items(), key=lambda x: -x[1]["ba"])]
    rows3.append(["All", sum(s["n"] for s in m["sp_stat"].values()), f"{m['tpa']:.1f}", f"{m['ba']:.1f}", f"{m['qmd']:.1f}", f"{m['cuft']:.0f}", f"{m['mbf']:.1f}"])
    t3 = Table(rows3, colWidths=[0.7 * inch, 0.5 * inch, 0.5 * inch, 0.5 * inch, 0.5 * inch, 0.55 * inch, 0.5 * inch]); t3.setStyle(tstyle(extra=[("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONTNAME", (0, -1), (-1, -1), FONT_B), ("LINEABOVE", (0, -1), (-1, -1), 0.4, colors.black)]))
    side = Table([[[Paragraph(f"TABLE 2. STAND TABLE BY {DBH_CLASS}-INCH DBH CLASS, PER ACRE", TT), t2], [Paragraph("TABLE 3. STOCK TABLE BY SPECIES, PER ACRE", TT), t3]]], colWidths=[3.6 * inch, 3.7 * inch])
    side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    page += [side, Paragraph("Trees: tally trees. TPA, BA (sq ft/ac), cubic feet and Scribner MBF per acre by BAF expansion. Species codes: PP ponderosa pine, WF white fir, DF Douglas-fir, SP sugar pine, IC incense-cedar.", FN)]
    page += [Spacer(1, 4), Image(mp, width=4.2 * inch, height=3.2 * inch), Paragraph("Figure 1. Plot locations with QA status.", FN), Spacer(1, 4)]
    fl = list(dict.fromkeys((p["plot"], f) for p in ps for f in flags.get(p["plot"], [])))
    page.append(Paragraph("<b>QA findings</b>: " + (f"{len(fl)} item(s) for crew follow-up" if fl else "none; unit accepted") + ("" if m["meets"] else f"; volume sampling error exceeds the {STRATUM_STD:.0f} % stratum standard" + ("" if m["binds"] else " (shown for information: unit under 10 % of sale volume)"))
                          + ("" if m["trees_ok"] else f"; under the R5 {MIN_MEASURED_TREES} measured trees per major species for " + ", ".join(m["trees_short"]) + " (sec. 41.3)") + ("" if m["practice_ok"] else f"; {PRACTICE_MIN_PLOTS - m['plots']} more plots to reach the R5 {PRACTICE_MIN_PLOTS}-plot minimum (sec. 41.3)"), B))
    if fl:
        ft = Table([["Plot", "Finding"]] + fl, colWidths=[0.9 * inch, 6.3 * inch]); ft.setStyle(tstyle(extra=[("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f6e0dc"))])); page.append(ft)
    rows = [["Plot", "Trees in", "BA sq ft/ac", "Slope %", "Status"]] + [[p["plot"], p["n_trees"], f"{p['n_trees'] * BAF:.0f}", p["slope"], "FLAG" if p["plot"] in flags else "ok"] for p in ps]
    pt = Table(rows, colWidths=[0.9 * inch, 0.9 * inch, 1.1 * inch, 0.8 * inch, 0.8 * inch], repeatRows=1); pt.setStyle(tstyle(extra=[("ALIGN", (1, 0), (3, -1), "RIGHT")]))
    page += [Spacer(1, 6), KeepTogether([Paragraph("TABLE 4. PLOT LIST", TT), pt])]
    return page


for u, g in units:
    page = build_page(u, g)
    if page is None:
        continue
    SimpleDocTemplate(os.path.join(OUT, f"Unit_{u['unit_id']}_Review.pdf"), pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=f"Unit {u['unit_id']} field data review", author="William Steinley").build(page, onFirstPage=watermark, onLaterPages=watermark)
    story += build_page(u, g) + [PageBreak()]      # fresh flowables: a built table cannot be reused

qa_rows = [["Unit", "Method", "Acres", "Plots", "cu ft/ac", "vol_cv", "vol_se95", "BA", "ba_cv", "ba_se95", "t", "Std %", "Design", "Share %", "Binds", "R5 20 trees", "R5 20-plot", "QA flags", "Status"]]
for u, g in units:
    m = M[u["unit_id"]]
    if not m:
        continue
    nf = len({(p["plot"], f) for p in plots if p["unit_id"] == u["unit_id"] for f in flags.get(p["plot"], [])})
    qa_rows.append([u["unit_id"], u["method"], f"{u['acres']:.0f}", m["plots"], f"{m['cuft']:.0f}", f"{m['vol_cv']:.0f}", f"{m['vol_se95']:.1f}", f"{m['ba']:.0f}", f"{m['ba_cv']:.0f}", f"{m['ba_se95']:.1f}", f"{m['tval']:.2f}", f"{m['se_std']:.0f}",
                    "meets" if m["meets"] else "fails", f"{m['share']:.1f}", "yes" if m["binds"] else "info", "ok" if m["trees_ok"] else "short", "ok" if m["practice_ok"] else "short", nf, "accepted" if (nf == 0 and m["meets"]) else "follow-up"])
qt = Table(qa_rows, repeatRows=1); qt.setStyle(tstyle(size=7.5, extra=[("ALIGN", (2, 0), (-2, -1), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)] + [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fbe9e2")) for i, r in enumerate(qa_rows) if i and r[-1] == "follow-up"]))
front = [Paragraph("Field Data Review, Sale-level QA Summary", H),
         Paragraph(f"Mohawk Valley West Slope demonstration. {len(plots)} plots, {len(trees)} trees, BAF {BAF:.0f} on a {GRID:.0f} ft grid across {len(units)} units ({tot_ac:,.0f} gross ac, {tot_net:,.0f} net ac after stream exclusion zones). "
                   "Plots tally all trees, so this is a pre-marking stand exam; the sale-cruise standards are applied as if the whole tallied volume were designated for cutting. "
                   f"Standards are those of Region 5 Supplement 2409.12-2021-4 to FSH 2409.12 ch. 40 (effective 2021-08-11), which replaces the national sec. 41.1 figures for a Plumas sale. "
                   f"<b>Sale as a whole</b> (stratified by unit, area weights, t = 2): net cubic volume {sale_vol:,.0f} cu ft/ac, sampling error {sale_vol_se95:.1f} % at 95 % confidence (basal area {sale_ba:.0f} sq ft/ac, {sale_ba_se95:.1f} %, reference only). "
                   f"R5 sec. 41.1 volume error standard {sale_std} % for a tree-measurement sale valued at about ${sale_value:,.0f} on net acres ({sale_mbf:,.0f} MBF at ${STUMPAGE_PER_MBF:.2f}/MBF, the Region 5 FY2025 average sold value, Forest Service Cut and Sold report; "
                   f"{sale_mbf_gross:,.0f} MBF, ${sale_value_gross:,.0f} on gross acres): <b>{'MEETS' if sale_meets else 'FAILS'}</b>. Basis: {sale_basis}; the national exhibit 01 tier for this value is {sale_std_national} %. "
                   f"<b>Strata</b>: each unit's volume sampling error is tested against the R5 sec. 41.1 stratum standard of {STRATUM_STD:.0f} % of volume for tree-measurement sales ({STRATUM_STD_SCALED:.0f} % scaled; national text 41.1(5)(b): {STRATUM_STD_NATIONAL:.0f} %). "
                   f"The standard binds for any stratum over {STRATUM_SHARE_PCT:.0f} % of total sale volume (Share %, Binds columns); smaller units are tested for information. {sum(1 for w, m, u in est if m['meets'])} of {len(est)} units meet it, worst {max(m['vol_se95'] for w, m, u in est):.1f} %; the standard binds for {sum(1 for w, m, u in est if m['binds'])} unit(s). The basal-area error is listed for reference. "
                   f"<b>Measured trees</b> (R5 sec. 41.3): at least {MIN_MEASURED_TREES} measured trees per major species, major being {MAJOR_SPECIES_PCT:.0f} % or more of sale value (one stumpage rate, so value share equals net-acre volume share): " + ", ".join(f"{k} {sale_sp_share[k]:.1f} %" for k in major_species) + ". "
                   f"Sale counts " + ", ".join(f"{k} {sale_measured[k]}" for k in SPECIES) + f": <b>{'MEETS' if sale_trees_ok else 'FAILS'}</b>; {sum(1 for w, m, u in est if m['trees_ok'])} of {len(est)} units meet it on their own (R5 20 trees column). "
                   f"<b>Plots</b>: the R5 {PRACTICE_MIN_PLOTS}-plot minimum for an area-based cruise (sec. 41.3); {sum(1 for w, m, u in est if m['practice_ok'])} of {len(est)} units have it (R5 20-plot column). "
                   f"QA flags are plots that failed a record check: {len(found)} flagged, {sum(1 for e in planted if e['plot'] in found)} of {len(planted)} planted errors caught.", B), Spacer(1, 8),
         Paragraph("TABLE A. CRUISE STATISTICS AND QA STATUS BY UNIT (STRATUM)", TT), qt,
         Paragraph("cu ft/ac: net cubic volume per acre (CVTS net of defect). vol_cv, ba_cv: coefficient of variation (%) of plot volume and plot basal area. vol_se95, ba_se95: sampling error (%), t x standard error / mean at 95 % confidence, Student's t on n-1 df. "
                   f"Std %: R5 stratum volume error standard. Design: vol_se95 against Std %. Share %: unit share of sale volume on net acres; Binds: yes where the share exceeds {STRATUM_SHARE_PCT:.0f} %, info otherwise. R5 20 trees: the R5 {MIN_MEASURED_TREES} measured trees per major species (by sale value), counted within the unit. R5 20-plot: the R5 {PRACTICE_MIN_PLOTS}-plot minimum (sec. 41.3). BA in sq ft/ac. "
                   f"Sources: {R5_SUPP_URL} ; {FSH_URL}", FN), PageBreak()]
merged.build(front + story, onFirstPage=watermark, onLaterPages=watermark)
csv_rows = [qa_rows[0] + ["data"]] + [r + ["simulated"] for r in qa_rows[1:]]      # data column marks every row as simulated
open(os.path.join(OUT, "qa_summary.csv"), "w", newline="").write("\n".join(",".join(str(c) for c in r) for r in csv_rows) + "\n")
print("review sheets written for", len(units), "units")
